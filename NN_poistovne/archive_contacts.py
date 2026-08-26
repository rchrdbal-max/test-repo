"""
Hard-delete every contact listed in the 'export final' sheet via the
ActiveCampaign API. Status is written to a new 'delete_status' column so the
script is resumable: re-runs skip rows whose delete_status is already set.

DESTRUCTIVE — DELETE /api/3/contacts/{id} is irreversible on this AC plan.
"""
from __future__ import annotations

import os
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).parent
EXCEL_PATH = ROOT / "Active_campaign-all data.xlsx"

load_dotenv(ROOT / ".env")
AC_API_URL = os.environ["AC_API_URL"].rstrip("/")
AC_API_KEY = os.environ["AC_API_KEY"]

SESSION = requests.Session()
SESSION.headers.update({"Api-Token": AC_API_KEY, "Accept": "application/json"})

DATA_SHEET = "export final"
ID_COL = 1  # column A
STATUS_HEADER = "delete_status"
SAVE_EVERY = 200
SLEEP_BETWEEN = 0.2  # 5 req/s pacing


def delete_contact(cid: str) -> tuple[bool, str]:
    """Returns (success, status_string)."""
    url = f"{AC_API_URL}/api/3/contacts/{cid}"
    last_err = ""
    for attempt in range(6):
        try:
            r = SESSION.delete(url, timeout=60)
        except (requests.exceptions.ReadTimeout,
                requests.exceptions.ConnectionError,
                requests.exceptions.ChunkedEncodingError) as e:
            last_err = f"{type(e).__name__}: {e}"
            time.sleep(min(2 ** attempt, 30))
            continue
        if r.status_code == 200:
            return True, f"DELETED {datetime.utcnow().isoformat(timespec='seconds')}Z"
        if r.status_code == 404:
            return True, "ALREADY_GONE"
        if r.status_code == 429:
            wait = int(r.headers.get("Retry-After", "5"))
            time.sleep(wait)
            continue
        if r.status_code >= 500:
            time.sleep(min(2 ** attempt, 30))
            last_err = f"HTTP {r.status_code}"
            continue
        # 4xx other than 404/429 — don't retry
        return False, f"ERROR HTTP {r.status_code}: {r.text[:120]}"
    return False, f"ERROR after retries: {last_err}"


def main():
    if not EXCEL_PATH.exists():
        print(f"ERROR: {EXCEL_PATH} not found", file=sys.stderr)
        sys.exit(1)

    print(f"Loading {EXCEL_PATH.name}...")
    wb = load_workbook(EXCEL_PATH)
    ws = wb[DATA_SHEET]

    headers = [c.value for c in ws[1]]
    if STATUS_HEADER in headers:
        status_col = headers.index(STATUS_HEADER) + 1
        print(f"Resuming — '{STATUS_HEADER}' already at col {status_col}")
    else:
        status_col = ws.max_column + 1
        ws.cell(row=1, column=status_col).value = STATUS_HEADER
        print(f"Created '{STATUS_HEADER}' at col {status_col}")

    total_rows = ws.max_row - 1
    todo: list[tuple[int, str]] = []
    skipped = 0
    for r in range(2, ws.max_row + 1):
        raw = ws.cell(row=r, column=ID_COL).value
        if raw is None:
            continue
        cid = str(int(raw)) if isinstance(raw, float) else str(raw)
        existing = ws.cell(row=r, column=status_col).value
        if existing:
            skipped += 1
            continue
        todo.append((r, cid))

    print(f"To delete: {len(todo):,}  (already-marked: {skipped:,}, total rows: {total_rows:,})")
    if not todo:
        print("Nothing to do.")
        return

    deleted = 0
    errors = 0
    t0 = time.time()
    for i, (row_idx, cid) in enumerate(todo, start=1):
        ok, status = delete_contact(cid)
        ws.cell(row=row_idx, column=status_col).value = status
        if ok:
            deleted += 1
        else:
            errors += 1
            print(f"  [row {row_idx} id={cid}] {status}", flush=True)

        if i % SAVE_EVERY == 0:
            wb.save(EXCEL_PATH)
            elapsed = time.time() - t0
            rate = i / elapsed if elapsed else 0
            eta = (len(todo) - i) / rate if rate else 0
            print(f"  progress {i}/{len(todo)}  deleted={deleted}  errors={errors}  "
                  f"({rate:.1f}/s, eta {eta/60:.1f}min)", flush=True)

        time.sleep(SLEEP_BETWEEN)

    wb.save(EXCEL_PATH)
    elapsed = time.time() - t0
    print(f"Done. deleted={deleted}  errors={errors}  elapsed={elapsed/60:.1f}min")


if __name__ == "__main__":
    main()
